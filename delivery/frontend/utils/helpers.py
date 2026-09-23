import io
import pandas as pd
import math
import markdown as md
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from utils.constants import settings


def confidence_label(score) -> str:
    if score is None:
        return "No score"
    try:
        score = float(score)
    except (ValueError, TypeError):
        return "No score"
    if math.isnan(score):
        return "No score"
    if score >= 0.85:
        return f"{round(score * 100)}% - high"
    if score >= settings.confidence_threshold:
        return f"{round(score * 100)}% - check"
    return f"{round(score * 100)}% - low"


def is_low_confidence(score) -> bool:
    if score is None:
        return True
    try:
        score = float(score)
    except (ValueError, TypeError):
        return True
    if math.isnan(score):
        return True
    return score < settings.confidence_threshold


def format_file_size(size_bytes: int) -> str:
    """
    Returns a human-readable file size string.
    Used on the upload screen alongside the filename
    """
    if size_bytes < 1024:
        return f"{size_bytes} B"
    if size_bytes < 1024 * 1024:
        return f"{round(size_bytes / 1024, 1)} KB"
    return f"{round(size_bytes / (1024 * 1024), 1)} MB"


def df_to_xlsx(df: pd.DataFrame) -> bytes:
    """
    Converts a pandas dataframe to an in-memory Excel file.
    Returns bytes suitable for st.download_button.
    """
    buf = io.BytesIO()
    df.to_excel(buf, index=False, engine="openpyxl")
    buf.seek(0)
    return buf.getvalue()



def render_markdown_small(text: str) -> str:
    """
    Converts markdown text to HTML wrapped in a small font container.
    Used for rendering extracted document text without oversized headings.
    """
    html_content = md.markdown(text)
    return f"""
        <div style="font-size: 0.8rem; line-height: 1.6;">
            {html_content}
        </div>
    """


def to_excel_bytes(df: pd.DataFrame) -> bytes:
    """
    Converts a dataframe to an in-memory .xlsx file and returns the raw bytes.
    Column headers are bolded. All columns are auto-sized to fit content.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracted Fields"

    header_font = Font(bold=True)
    header_fill = PatternFill(
        start_color="E4E4E4", end_color="E4E4E4", fill_type="solid"
    )

    for row_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
        ws.append(row)
        if row_idx == 1:
            for cell in ws[row_idx]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")

    # Auto-size columns to fit content
    for col in ws.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in col
        )
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 60)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


def is_low_confidence(score: float) -> bool:
    """
    Returns True if the score is below the configured confidence threshold.
    Returns True for None or NaN scores so they are flagged for review.
    """
    if score is None or (isinstance(score, float) and math.isnan(score)):
        return True
    return score < settings.confidence_threshold